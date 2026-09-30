import streamlit as st
import requests
from io import BytesIO

from docx import Document
from docx.shared import Pt

from fpdf import FPDF


BACKEND_URL = "http://127.0.0.1:8000/generate"


st.set_page_config(
    page_title="LegalEase",
    page_icon="⚖️",
    layout="wide"
)


st.title("⚖️ LegalEase")
st.subheader("AI-Powered Legal Document Generator")

st.write(
    "Create, edit and download professional legal documents "
    "using AI."
)


# -----------------------------
# INPUT SECTION
# -----------------------------

st.header("Document Details")

document_type = st.text_input(
    "Document Type",
    placeholder="Example: Employment Agreement"
)

parties = st.text_area(
    "Parties Involved",
    placeholder="Example: ABC Company and Employee"
)

terms = st.text_area(
    "Terms & Conditions",
    placeholder=(
        "Example: Job role; Salary; Working hours; "
        "Leave policy; Confidentiality"
    )
)

dates = st.text_input(
    "Effective Date",
    placeholder="Example: 01-10-2026"
)


# -----------------------------
# GENERATE BUTTON
# -----------------------------

if st.button("Generate Document", type="primary"):

    if not document_type:
        st.warning("Please enter Document Type.")

    elif not parties:
        st.warning("Please enter Parties.")

    elif not terms:
        st.warning("Please enter Terms.")

    elif not dates:
        st.warning("Please enter Effective Date.")

    else:

        payload = {
            "document_type": document_type,
            "parties": parties,
            "terms": terms,
            "dates": dates
        }

        try:

            with st.spinner("Generating document with Gemini AI..."):

                response = requests.post(
                    BACKEND_URL,
                    json=payload,
                    timeout=120
                )

            if response.status_code == 200:

                result = response.json()

                st.session_state["document"] = result["document"]

                st.success("Document generated successfully!")

            else:

                st.error(
                    f"Backend error: {response.status_code}"
                )

        except requests.exceptions.ConnectionError:

            st.error(
                "Cannot connect to FastAPI backend. "
                "Please make sure Uvicorn is running."
            )

        except Exception as e:

            st.error(f"Error: {e}")


# -----------------------------
# DOCUMENT PREVIEW / EDIT
# -----------------------------

if "document" in st.session_state:

    st.header("Document Preview")

    edited_document = st.text_area(
        "Edit Document",
        value=st.session_state["document"],
        height=500
    )

    st.session_state["document"] = edited_document


    # -----------------------------
    # TXT DOWNLOAD
    # -----------------------------

    txt_data = edited_document.encode("utf-8")

    st.download_button(
        label="Download TXT",
        data=txt_data,
        file_name="LegalEase_Document.txt",
        mime="text/plain"
    )


    # -----------------------------
    # DOCX DOWNLOAD
    # -----------------------------

    doc = Document()

    title = doc.add_paragraph()
    title.alignment = 1

    run = title.add_run(document_type)

    run.bold = True
    run.font.size = Pt(18)

    for paragraph in edited_document.split("\n"):

        if paragraph.strip():

            p = doc.add_paragraph(paragraph)

            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)


    docx_buffer = BytesIO()

    doc.save(docx_buffer)

    docx_buffer.seek(0)


    st.download_button(
        label="Download DOCX",
        data=docx_buffer,
        file_name="LegalEase_Document.docx",
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        )
    )
# -----------------------------
# PDF DOWNLOAD
# -----------------------------

pdf = FPDF()

pdf.set_auto_page_break(
    auto=True,
    margin=15
)

pdf.add_page()

pdf.set_font(
    "Arial",
    "B",
    16
)

pdf.cell(
    0,
    10,
    document_type,
    ln=True,
    align="C"
)

pdf.ln(5)

pdf.set_font(
    "Arial",
    size=11
)

safe_text = edited_document.encode(
    "latin-1",
    "replace"
).decode(
    "latin-1"
)

for line in safe_text.split("\n"):

    if line.strip():

        pdf.multi_cell(
            0,
            7,
            line
        )

        pdf.ln(1)

pdf_output = pdf.output(dest="S")

if isinstance(pdf_output, str):
    pdf_bytes = pdf_output.encode("latin-1")
else:
    pdf_bytes = bytes(pdf_output)

st.download_button(
    label="Download PDF",
    data=pdf_bytes,
    file_name="LegalEase_Document.pdf",
    mime="application/pdf"
)